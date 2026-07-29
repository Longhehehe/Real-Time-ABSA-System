"""Build a label-blind, deduplicated re-annotation package for the old XLSX set.

The legacy workbooks are immutable source evidence.  Their nine historical
aspect columns are never copied into the clean-core records or LLM target
inputs.  Only a one-way hash of each old label vector is retained in the
source-row decision ledger so that the stripping decision can be audited
without exposing the old labels to the annotation run.
"""

from __future__ import annotations

import argparse
from collections import Counter
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import shutil
import tempfile
import unicodedata
from typing import Any, Iterable


REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
SOURCE_ROOT = REPOSITORY_ROOT / "src"

EXPECTED_HEADERS = (
    "reviewContent",
    "Chất lượng sản phẩm",
    "Hiệu năng & Trải nghiệm",
    "Đúng mô tả",
    "Giá cả & Khuyến mãi",
    "Vận chuyển",
    "Đóng gói",
    "Dịch vụ & Thái độ Shop",
    "Bảo hành & Đổi trả",
    "Tính xác thực",
)
LABEL_HEADERS = EXPECTED_HEADERS[1:]
INPUT_SCHEMA_VERSION = "absa-ai-tranche-input/1.0.0"
PRIVATE_SCHEMA_VERSION = "absa-ai-tranche-private-index/1.0.0"
SOURCE_SCHEMA_VERSION = "legacy-old-label-blind-clean-core/1.0.0"
SOURCE_LEDGER_SCHEMA_VERSION = "legacy-old-source-row-ledger/1.0.0"
SELECTION_LEDGER_SCHEMA_VERSION = "absa-pseudolabel-selection-ledger/1.0.0"
SELECTION_SPEC_VERSION = "absa-legacy-old-relabel-selection/1.0.0"
TRANCHE_NAME = "tranche-0003"
SELECTION_DECISION = "SELECT_TRANCHE_0003"

DEFAULT_SOURCE_DIR = Path("legacy/data/old_dataset")
DEFAULT_SOURCE_RELEASE = Path(
    "data/releases/legacy_old_reviews_label_blind_v1_20260728"
)
DEFAULT_PACKAGE = Path(
    "data/annotations/absa_legacy_old_relabel_9772_v1_20260728"
)
DEFAULT_TEMPLATE_PACKAGE = Path(
    "data/annotations/absa_ai_remainder_8976_v1_20260728"
)
DEFAULT_GUIDELINE = Path("docs/ABSA_ANNOTATION_GUIDELINE_V2.md")


def _canonical_json(value: Any) -> str:
    return json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    )


def _sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def _sha256_text(value: str) -> str:
    return _sha256_bytes(value.encode("utf-8"))


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _natural_part(path: Path) -> tuple[int, str]:
    match = re.search(r"part(\d+)", path.stem, flags=re.IGNORECASE)
    return (int(match.group(1)) if match else 10**9, path.name.casefold())


def _normalized_key(text: str) -> str:
    return re.sub(
        r"\s+",
        " ",
        unicodedata.normalize("NFKC", text).casefold(),
    ).strip()


def _write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
        newline="\n",
    )


def _write_jsonl(path: Path, rows: Iterable[dict[str, Any]]) -> int:
    path.parent.mkdir(parents=True, exist_ok=True)
    count = 0
    with path.open("w", encoding="utf-8", newline="\n") as handle:
        for row in rows:
            handle.write(_canonical_json(row) + "\n")
            count += 1
    return count


def _artifact(
    path: Path,
    root: Path,
    *,
    records: int | None = None,
) -> dict[str, Any]:
    value: dict[str, Any] = {
        "path": path.relative_to(root).as_posix(),
        "bytes": path.stat().st_size,
        "sha256": _sha256_file(path),
    }
    if records is not None:
        value["records"] = records
    return value


def _write_checksum_closure(
    root: Path,
    artifacts: list[dict[str, Any]],
    *,
    manifest_name: str,
    sums_name: str,
) -> None:
    rows = [(item["sha256"], item["path"]) for item in artifacts]
    rows.append((_sha256_file(root / manifest_name), manifest_name))
    (root / sums_name).write_text(
        "".join(
            f"{digest}  {relative}\n"
            for digest, relative in sorted(rows, key=lambda item: item[1])
        ),
        encoding="utf-8",
        newline="\n",
    )


def _stable_id(prefix: str, *parts: str, length: int = 20) -> str:
    digest = _sha256_text("\x1f".join(parts))
    return f"{prefix}{digest[:length]}"


def _load_workbooks(source_dir: Path) -> tuple[
    list[dict[str, Any]],
    list[dict[str, Any]],
    list[dict[str, Any]],
]:
    try:
        import openpyxl
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "openpyxl is required only for this legacy XLSX import. "
            "Use the system Python environment recorded by the legacy audit."
        ) from exc

    files = sorted(source_dir.glob("*.xlsx"), key=_natural_part)
    if len(files) != 10:
        raise ValueError(f"Expected exactly 10 old workbooks, got {len(files)}")

    source_rows: list[dict[str, Any]] = []
    inventory: list[dict[str, Any]] = []
    normalized_groups: dict[str, list[int]] = {}
    for workbook_rank, path in enumerate(files, 1):
        workbook_sha = _sha256_file(path)
        workbook = openpyxl.load_workbook(
            path,
            read_only=True,
            data_only=True,
        )
        worksheet = workbook.active
        headers = tuple(
            cell.value
            for cell in next(
                worksheet.iter_rows(min_row=1, max_row=1)
            )
        )
        if headers != EXPECTED_HEADERS:
            workbook.close()
            raise ValueError(f"Unexpected headers in {path}: {headers!r}")
        valid_rows = 0
        for row_number, values in enumerate(
            worksheet.iter_rows(min_row=2, values_only=True),
            2,
        ):
            raw_review = values[0]
            if raw_review is None or not str(raw_review).strip():
                continue
            valid_rows += 1
            original_text = str(raw_review)
            curated_text = original_text.strip()
            normalized = _normalized_key(original_text)
            if not normalized:
                raise ValueError(f"Normalized review is empty: {path}:{row_number}")
            old_label_values = list(values[1:])
            old_label_vector_sha = _sha256_text(
                _canonical_json(old_label_values)
            )
            row = {
                "source_file": path.name,
                "source_file_sha256": workbook_sha,
                "source_sheet": worksheet.title,
                "source_row_number": row_number,
                "source_position": f"{path.name}:{worksheet.title}:{row_number}",
                "source_review_text": original_text,
                "curated_review_text": curated_text,
                "source_review_text_sha256": _sha256_text(original_text),
                "curated_review_text_sha256": _sha256_text(curated_text),
                "normalized_review_key": normalized,
                "normalized_review_key_sha256": _sha256_text(normalized),
                "legacy_label_vector_sha256": old_label_vector_sha,
                "legacy_label_cells": len(old_label_values),
                "legacy_label_nonempty_cells": sum(
                    value is not None and str(value).strip() != ""
                    for value in old_label_values
                ),
                "text_was_trimmed": curated_text != original_text,
                "workbook_rank": workbook_rank,
            }
            source_rows.append(row)
            normalized_groups.setdefault(normalized, []).append(
                len(source_rows) - 1
            )
        inventory.append(
            {
                "path": path.relative_to(REPOSITORY_ROOT).as_posix(),
                "bytes": path.stat().st_size,
                "sha256": workbook_sha,
                "worksheet": worksheet.title,
                "max_row": worksheet.max_row,
                "max_column": worksheet.max_column,
                "valid_review_rows": valid_rows,
                "headers": list(headers),
            }
        )
        workbook.close()

    clean_rows: list[dict[str, Any]] = []
    source_ledger: list[dict[str, Any]] = []
    canonical_by_normalized: dict[str, dict[str, Any]] = {}
    for normalized, positions in normalized_groups.items():
        representative = source_rows[positions[0]]
        sample_id = _stable_id(
            "legacy-old-",
            representative["normalized_review_key_sha256"],
            length=24,
        )
        canonical = {
            "schema_version": SOURCE_SCHEMA_VERSION,
            "sample_id": sample_id,
            "curated_review_text": representative["curated_review_text"],
            "category": "",
            "rating": None,
            "collection_transport": "legacy_xlsx",
            "product_id": None,
            "curation": {
                "status": (
                    "KEEP_CLEANED"
                    if representative["text_was_trimmed"]
                    else "KEEP"
                ),
                "parent_canonical_row": representative["source_position"],
                "curated_text_sha256": representative[
                    "curated_review_text_sha256"
                ],
                "deduplication_key": "NFKC_CASEFOLD_WHITESPACE",
                "normalized_review_key_sha256": representative[
                    "normalized_review_key_sha256"
                ],
                "source_alias_count": len(positions),
                "legacy_labels_removed": True,
            },
        }
        canonical_by_normalized[normalized] = canonical
        clean_rows.append(canonical)

    canonical_rank = {
        row["sample_id"]: rank for rank, row in enumerate(clean_rows, 1)
    }
    for row in source_rows:
        canonical = canonical_by_normalized[row["normalized_review_key"]]
        is_representative = (
            row["source_position"]
            == canonical["curation"]["parent_canonical_row"]
        )
        source_ledger.append(
            {
                "schema_version": SOURCE_LEDGER_SCHEMA_VERSION,
                "source_file": row["source_file"],
                "source_file_sha256": row["source_file_sha256"],
                "source_sheet": row["source_sheet"],
                "source_row_number": row["source_row_number"],
                "source_position": row["source_position"],
                "source_review_text_sha256": row[
                    "source_review_text_sha256"
                ],
                "curated_review_text_sha256": row[
                    "curated_review_text_sha256"
                ],
                "normalized_review_key_sha256": row[
                    "normalized_review_key_sha256"
                ],
                "canonical_sample_id": canonical["sample_id"],
                "canonical_rank": canonical_rank[canonical["sample_id"]],
                "decision": (
                    "CANONICAL_REPRESENTATIVE_LABELS_STRIPPED"
                    if is_representative
                    else "DUPLICATE_ALIAS_MAP_TO_CANONICAL"
                ),
                "legacy_label_columns_removed_from_llm_input": list(
                    LABEL_HEADERS
                ),
                "legacy_label_cells": row["legacy_label_cells"],
                "legacy_label_nonempty_cells": row[
                    "legacy_label_nonempty_cells"
                ],
                "legacy_label_vector_sha256": row[
                    "legacy_label_vector_sha256"
                ],
                "legacy_label_values_copied_to_llm_input": False,
            }
        )
    return clean_rows, source_ledger, inventory


def _build_source_release(
    *,
    source_dir: Path,
    output: Path,
) -> dict[str, Any]:
    if output.exists():
        raise FileExistsError(f"Source release already exists: {output}")
    clean_rows, source_ledger, inventory = _load_workbooks(source_dir)
    if len(source_ledger) != 10105 or len(clean_rows) != 9772:
        raise ValueError(
            "Frozen legacy inventory mismatch: expected 10105 source rows "
            f"and 9772 normalized-unique rows, got {len(source_ledger)} and "
            f"{len(clean_rows)}"
        )
    text_hashes = {
        row["curation"]["curated_text_sha256"] for row in clean_rows
    }
    if len(text_hashes) != len(clean_rows):
        raise ValueError("Clean core contains duplicate target text hashes")

    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = Path(
        tempfile.mkdtemp(prefix=f".{output.name}.", dir=output.parent)
    )
    try:
        clean_path = temporary / "clean_core.jsonl"
        ledger_path = temporary / "source_row_decision_ledger.jsonl"
        inventory_path = temporary / "source_workbook_inventory.json"
        scrub_path = temporary / "legacy_label_scrub_audit.json"
        _write_jsonl(clean_path, clean_rows)
        _write_jsonl(ledger_path, source_ledger)
        _write_json(inventory_path, {"workbooks": inventory})
        _write_json(
            scrub_path,
            {
                "schema_version": "legacy-label-scrub-audit/1.0.0",
                "status": "PASS",
                "source_rows": len(source_ledger),
                "canonical_unique_reviews": len(clean_rows),
                "duplicate_alias_rows": len(source_ledger) - len(clean_rows),
                "historical_label_columns": list(LABEL_HEADERS),
                "historical_label_column_count": len(LABEL_HEADERS),
                "historical_label_values_in_clean_core": 0,
                "historical_label_values_in_llm_target_input": 0,
                "source_workbooks_modified": False,
                "policy": (
                    "Historical label values are omitted from clean_core and "
                    "LLM targets. Source workbooks remain immutable evidence; "
                    "only one-way label-vector hashes survive in the ledger."
                ),
            },
        )
        artifacts = [
            _artifact(clean_path, temporary, records=len(clean_rows)),
            _artifact(
                ledger_path,
                temporary,
                records=len(source_ledger),
            ),
            _artifact(inventory_path, temporary),
            _artifact(scrub_path, temporary),
        ]
        release_id = _stable_id(
            "legacy-old-label-blind-",
            _sha256_file(clean_path),
            _sha256_file(ledger_path),
            length=16,
        )
        manifest = {
            "schema_version": "legacy-old-source-release-manifest/1.0.0",
            "artifact_type": "LEGACY_OLD_LABEL_BLIND_CANONICAL_SOURCE",
            "status": "FROZEN_LABELS_STRIPPED_FROM_TARGETS",
            "release_id": release_id,
            "built_at": datetime.now(timezone.utc).isoformat(),
            "source_directory": source_dir.as_posix(),
            "source_rows": len(source_ledger),
            "clean_core_records": len(clean_rows),
            "duplicate_alias_rows": len(source_ledger) - len(clean_rows),
            "deduplication": {
                "key": "NFKC_CASEFOLD_WHITESPACE",
                "representative": "first row in part-number then worksheet order",
                "labels_propagated_to_aliases": False,
            },
            "legacy_label_handling": {
                "columns": list(LABEL_HEADERS),
                "values_copied_to_clean_core": False,
                "values_copied_to_llm_targets": False,
                "source_workbooks_modified": False,
                "retained_provenance": "one-way label-vector SHA-256 only",
            },
            "artifacts": sorted(artifacts, key=lambda item: item["path"]),
        }
        manifest_path = temporary / "manifest.json"
        _write_json(manifest_path, manifest)
        _write_checksum_closure(
            temporary,
            artifacts,
            manifest_name="manifest.json",
            sums_name="SHA256SUMS.txt",
        )
        temporary.replace(output)
    except Exception:
        shutil.rmtree(temporary, ignore_errors=True)
        raise
    return manifest


def _load_json(path: Path) -> dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError(f"Expected JSON object: {path}")
    return value


def _load_jsonl(path: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    with path.open("r", encoding="utf-8") as handle:
        for line_number, line in enumerate(handle, 1):
            if not line.strip():
                continue
            value = json.loads(line)
            if not isinstance(value, dict):
                raise ValueError(f"Expected object at {path}:{line_number}")
            rows.append(value)
    return rows


def _build_package(
    *,
    source_release: Path,
    output: Path,
    template_package: Path,
    guideline: Path,
) -> dict[str, Any]:
    if output.exists():
        raise FileExistsError(f"Annotation package already exists: {output}")
    source_manifest = _load_json(source_release / "manifest.json")
    template_manifest = _load_json(template_package / "prepare_manifest.json")
    clean_rows = _load_jsonl(source_release / "clean_core.jsonl")
    calibration_rows = _load_jsonl(
        template_package / "calibration" / "human_confirmed.jsonl"
    )
    diagnostic_split = _load_json(
        template_package / "calibration" / "diagnostic_split.json"
    )
    if len(clean_rows) != 9772:
        raise ValueError(f"Expected 9772 legacy clean rows, got {len(clean_rows)}")
    if _sha256_file(guideline) != template_manifest["guideline"]["sha256"]:
        raise ValueError("Guideline differs from the frozen prior configuration")

    selection = template_manifest["selection"]
    human_manifest_path = Path(selection["human_reference_manifest_path"])
    reservation_path = Path(selection["group_reservations_path"])
    reservation_rows = _load_jsonl(reservation_path)
    reserved_sample_ids = {row["sample_id"] for row in reservation_rows}
    reserved_text_hashes = {
        row["review_text_sha256"] for row in reservation_rows
    }
    target_sample_ids = {row["sample_id"] for row in clean_rows}
    target_text_hashes = {
        row["curation"]["curated_text_sha256"] for row in clean_rows
    }
    sample_overlap = target_sample_ids.intersection(reserved_sample_ids)
    text_overlap = target_text_hashes.intersection(reserved_text_hashes)
    if sample_overlap:
        raise ValueError("Legacy sample IDs overlap reserved sample IDs")

    source_manifest_sha = _sha256_file(source_release / "manifest.json")
    clean_core_sha = _sha256_file(source_release / "clean_core.jsonl")
    tranche_id = _stable_id(
        "absa-ai-tranche-",
        source_manifest["release_id"],
        SELECTION_SPEC_VERSION,
        TRANCHE_NAME,
        source_manifest_sha,
        clean_core_sha,
        template_manifest["human_calibration"]["calibration_payload_sha256"],
        length=16,
    )

    blind_rows: list[dict[str, Any]] = []
    private_rows: list[dict[str, Any]] = []
    selection_rows: list[dict[str, Any]] = []
    selected_status: Counter[str] = Counter()
    for rank, source in enumerate(clean_rows, 1):
        text = source["curated_review_text"]
        text_sha = source["curation"]["curated_text_sha256"]
        annotation_id = _stable_id(
            "aold-",
            tranche_id,
            source["sample_id"],
            length=20,
        )
        blind_rows.append(
            {
                "schema_version": INPUT_SCHEMA_VERSION,
                "selection_rank": rank,
                "annotation_id": annotation_id,
                "reviewContent": text,
                "review_text_sha256": text_sha,
            }
        )
        private_rows.append(
            {
                "schema_version": PRIVATE_SCHEMA_VERSION,
                "selection_rank": rank,
                "annotation_id": annotation_id,
                "sample_id": source["sample_id"],
                "review_text_sha256": text_sha,
                "parent_canonical_row": source["curation"][
                    "parent_canonical_row"
                ],
                "curation_status": source["curation"]["status"],
                "category": source.get("category", ""),
                "rating": source.get("rating"),
                "collection_transport": source.get("collection_transport"),
                "product_id": source.get("product_id"),
                "source_release_id": source_manifest["release_id"],
            }
        )
        selection_rows.append(
            {
                "schema_version": SELECTION_LEDGER_SCHEMA_VERSION,
                "sample_id": source["sample_id"],
                "review_text_sha256": text_sha,
                "rating": source.get("rating"),
                "category": "<blank>",
                "collection_transport": "legacy_xlsx",
                "selection_rank_digest": _sha256_text(
                    "\x1f".join(
                        (
                            source_manifest["release_id"],
                            source["sample_id"],
                            text_sha,
                            TRANCHE_NAME,
                        )
                    )
                ),
                "selected": True,
                "selection_rank": rank,
                "decision": SELECTION_DECISION,
            }
        )
        selected_status[source["curation"]["status"]] += 1

    allowed_blind_keys = {
        "schema_version",
        "selection_rank",
        "annotation_id",
        "reviewContent",
        "review_text_sha256",
    }
    if any(set(row) != allowed_blind_keys for row in blind_rows):
        raise ValueError("Blind targets contain unexpected fields")
    if any(header in _canonical_json(blind_rows) for header in LABEL_HEADERS):
        raise ValueError("Historical label column leaked into blind targets")

    membership_sha = _sha256_text(
        "".join(
            f"{row['sample_id']}\t{row['review_text_sha256']}\n"
            for row in private_rows
        )
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = Path(
        tempfile.mkdtemp(prefix=f".{output.name}.", dir=output.parent)
    )
    try:
        blind_path = temporary / "input" / "blind_reviews.jsonl"
        private_path = temporary / "input" / "private_index.jsonl"
        selection_path = temporary / "input" / "selection_ledger.jsonl"
        calibration_path = (
            temporary / "calibration" / "human_confirmed.jsonl"
        )
        split_path = temporary / "calibration" / "diagnostic_split.json"
        _write_jsonl(blind_path, blind_rows)
        _write_jsonl(private_path, private_rows)
        _write_jsonl(selection_path, selection_rows)
        _write_jsonl(calibration_path, calibration_rows)
        _write_json(split_path, diagnostic_split)

        provenance = temporary / "provenance"
        provenance.mkdir(parents=True, exist_ok=True)
        provenance_sources = {
            guideline: provenance / "ABSA_ANNOTATION_GUIDELINE_V2.md",
            Path(__file__).resolve(): (
                provenance / "prepare_legacy_old_reannotation.py"
            ),
            REPOSITORY_ROOT / "scripts" / "run_ai_annotation_tranche.py": (
                provenance / "run_ai_annotation_tranche.py"
            ),
            SOURCE_ROOT / "lazada_collector" / "ai_tranche.py": (
                provenance / "ai_tranche.py"
            ),
            SOURCE_ROOT / "lazada_collector" / "llm_annotation.py": (
                provenance / "llm_annotation.py"
            ),
            SOURCE_ROOT / "lazada_collector" / "llm_backends.py": (
                provenance / "llm_backends.py"
            ),
            REPOSITORY_ROOT
            / "configs"
            / "absa_compact_batch_output_schema_v1.json": (
                provenance / "absa_compact_batch_output_schema_v1.json"
            ),
        }
        for source, destination in provenance_sources.items():
            if not source.is_file():
                raise FileNotFoundError(source)
            shutil.copy2(source, destination)

        artifacts = [
            _artifact(blind_path, temporary, records=len(blind_rows)),
            _artifact(private_path, temporary, records=len(private_rows)),
            _artifact(
                selection_path,
                temporary,
                records=len(selection_rows),
            ),
            _artifact(
                calibration_path,
                temporary,
                records=len(calibration_rows),
            ),
            _artifact(split_path, temporary),
        ]
        artifacts.extend(
            _artifact(path, temporary)
            for path in sorted(provenance.iterdir())
            if path.is_file()
        )
        human_calibration = dict(template_manifest["human_calibration"])
        manifest = {
            "artifact_type": "CALIBRATED_ABSA_AI_TRANCHE_WORK_PACKAGE",
            "status": "PREPARED_NOT_LABELED",
            "tranche_id": tranche_id,
            "built_at": datetime.now(timezone.utc).isoformat(),
            "target_records": len(blind_rows),
            "selection": {
                "selection_spec_version": SELECTION_SPEC_VERSION,
                "tranche": TRANCHE_NAME,
                "method": (
                    "Select every normalized-unique legacy old-dataset review "
                    "in deterministic source order after stripping all nine "
                    "historical aspect-label columns. Duplicate source rows "
                    "remain in the source ledger and map to one canonical "
                    "target."
                ),
                "stratum_order": ["legacy_source_order"],
                "stratum_allocations": {
                    "legacy_xlsx": {
                        "available": len(blind_rows),
                        "allocated": len(blind_rows),
                    }
                },
                "ordered_membership_sha256": membership_sha,
                "ordered_membership_serialization": (
                    "UTF-8 sample_id<TAB>review_text_sha256<LF> repeated in "
                    "selection_rank order, including terminal LF"
                ),
                "human_reference_manifest_path": human_manifest_path.as_posix(),
                "human_reference_manifest_sha256": selection[
                    "human_reference_manifest_sha256"
                ],
                "eligible_after_reference_exclusion": len(blind_rows),
                "eligible_after_prior_tranche_exclusion": len(blind_rows),
                "excluded_prior_tranche_records": 0,
                "excluded_prior_tranche": None,
                "reserved_human_reference_records": selection[
                    "reserved_human_reference_records"
                ],
                "reserved_reference_group_records": selection[
                    "reserved_reference_group_records"
                ],
                "group_reservations_path": reservation_path.as_posix(),
                "group_reservations_sha256": selection[
                    "group_reservations_sha256"
                ],
                "overlap_with_human_reference": 0,
                "overlap_with_reserved_reference_groups": 0,
                "overlap_with_prior_tranche": 0,
                "cross_corpus_reserved_text_hash_overlap": len(text_overlap),
                "cross_corpus_reserved_text_hash_overlap_policy": (
                    "Flag in provenance; sample IDs remain disjoint. These "
                    "records must not be used as independent benchmark items."
                ),
            },
            "source_release": {
                "path": source_release.as_posix(),
                "release_id": source_manifest["release_id"],
                "manifest_sha256": source_manifest_sha,
                "clean_core_sha256": clean_core_sha,
                "clean_core_records": len(clean_rows),
            },
            "guideline": {
                "path": guideline.as_posix(),
                "version": template_manifest["guideline"]["version"],
                "sha256": _sha256_file(guideline),
            },
            "human_calibration": human_calibration,
            "legacy_label_handling": {
                "source_columns": list(LABEL_HEADERS),
                "old_label_values_in_blind_targets": 0,
                "old_label_values_used_for_selection": False,
                "old_label_values_used_for_prompt_calibration": False,
                "old_label_values_used_for_validation": False,
                "source_workbooks_modified": False,
                "source_release_scrub_audit": (
                    source_release / "legacy_label_scrub_audit.json"
                ).as_posix(),
            },
            "distributions": {
                "source_categories": {"": len(clean_rows)},
                "selected_categories": {"": len(clean_rows)},
                "selected_curation_status": dict(
                    sorted(selected_status.items())
                ),
            },
            "data_transmission_notice": (
                "Only blinded annotation_id and review text are sent to the "
                "configured LLM backend. Historical labels and private source "
                "positions are never included in the LLM request."
            ),
            "artifacts": sorted(artifacts, key=lambda item: item["path"]),
        }
        manifest_path = temporary / "prepare_manifest.json"
        _write_json(manifest_path, manifest)
        _write_checksum_closure(
            temporary,
            artifacts,
            manifest_name="prepare_manifest.json",
            sums_name="INPUT_SHA256SUMS.txt",
        )
        temporary.replace(output)
    except Exception:
        shutil.rmtree(temporary, ignore_errors=True)
        raise
    return manifest


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-dir", type=Path, default=DEFAULT_SOURCE_DIR)
    parser.add_argument(
        "--source-release",
        type=Path,
        default=DEFAULT_SOURCE_RELEASE,
    )
    parser.add_argument("--output", type=Path, default=DEFAULT_PACKAGE)
    parser.add_argument(
        "--template-package",
        type=Path,
        default=DEFAULT_TEMPLATE_PACKAGE,
    )
    parser.add_argument("--guideline", type=Path, default=DEFAULT_GUIDELINE)
    return parser


def main() -> None:
    args = build_parser().parse_args()
    source_dir = args.source_dir.resolve()
    source_release = args.source_release.resolve()
    output = args.output.resolve()
    template_package = args.template_package.resolve()
    guideline = args.guideline.resolve()
    source_manifest = _build_source_release(
        source_dir=source_dir,
        output=source_release,
    )
    package_manifest = _build_package(
        source_release=source_release,
        output=output,
        template_package=template_package,
        guideline=guideline,
    )
    print(
        json.dumps(
            {
                "status": "PREPARED_NOT_LABELED",
                "source_release": str(source_release),
                "source_release_id": source_manifest["release_id"],
                "source_rows": source_manifest["source_rows"],
                "target_records": package_manifest["target_records"],
                "duplicate_alias_rows": source_manifest[
                    "duplicate_alias_rows"
                ],
                "old_label_values_in_blind_targets": 0,
                "package": str(output),
                "tranche_id": package_manifest["tranche_id"],
                "reserved_text_hash_overlap": package_manifest["selection"][
                    "cross_corpus_reserved_text_hash_overlap"
                ],
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
