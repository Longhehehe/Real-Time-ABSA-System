"""Independently validate the legacy-old ABSA re-annotation release."""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import sys
from typing import Any, Mapping


REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
SOURCE_ROOT = REPOSITORY_ROOT / "src"
if str(REPOSITORY_ROOT) not in sys.path:
    sys.path.insert(0, str(REPOSITORY_ROOT))
if str(SOURCE_ROOT) not in sys.path:
    sys.path.insert(0, str(SOURCE_ROOT))

from scripts.validate_ai_annotation_tranche import validate_release


ASPECTS = (
    "Chất lượng sản phẩm",
    "Hiệu năng & Trải nghiệm",
    "Đúng mô tả",
    "Giá cả & Khuyến mãi",
    "Vận chuyển",
    "Đóng gói",
    "Dịch vụ & Thái độ Shop",
    "Bảo hành & Đổi trả",
    "Tính xác thực",
)
EXPECTED_HEADERS = ("reviewContent", *ASPECTS)
DEFAULT_PACKAGE = Path(
    "data/annotations/absa_legacy_old_relabel_9772_v1_20260728"
)
DEFAULT_SOURCE_RELEASE = Path(
    "data/releases/legacy_old_reviews_label_blind_v1_20260728"
)
DEFAULT_SOURCE_DIR = Path("legacy/data/old_dataset")
DEFAULT_REFERENCE_PACKAGE = Path(
    "data/annotations/absa_ai_remainder_8976_v1_20260728"
)
DEFAULT_REPORT = Path(
    "docs/audits/legacy_old_reannotation_validation_20260728.json"
)


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _read_json(path: Path) -> dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError(f"Expected JSON object: {path}")
    return value


def _read_jsonl(path: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    with path.open("r", encoding="utf-8") as handle:
        for line_number, line in enumerate(handle, 1):
            if not line.strip():
                continue
            value = json.loads(line)
            if not isinstance(value, dict):
                raise ValueError(f"Expected object at {path}:{line_number}")
            rows.append(value)
    return rows


def _read_sums(path: Path) -> dict[str, str]:
    rows: dict[str, str] = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line:
            continue
        digest, relative = line.split("  ", 1)
        if relative in rows:
            raise ValueError(f"Duplicate checksum entry: {relative}")
        rows[relative] = digest
    return rows


def _validate_closure(
    root: Path,
    *,
    manifest_name: str = "manifest.json",
    sums_name: str = "SHA256SUMS.txt",
) -> dict[str, Any]:
    manifest = _read_json(root / manifest_name)
    artifacts = manifest.get("artifacts")
    if not isinstance(artifacts, list):
        raise ValueError(f"Malformed artifact inventory: {root}")
    expected = {manifest_name}
    sums = _read_sums(root / sums_name)
    if sums.get(manifest_name) != _sha256_file(root / manifest_name):
        raise ValueError(f"Manifest checksum mismatch: {root}")
    for item in artifacts:
        if not isinstance(item, Mapping):
            raise ValueError(f"Malformed artifact entry: {root}")
        relative = item.get("path")
        if not isinstance(relative, str):
            raise ValueError(f"Malformed artifact path: {root}")
        path = root / relative
        if (
            not path.is_file()
            or item.get("bytes") != path.stat().st_size
            or item.get("sha256") != _sha256_file(path)
            or sums.get(relative) != _sha256_file(path)
        ):
            raise ValueError(f"Artifact mismatch: {root / relative}")
        expected.add(relative)
    if set(sums) != expected:
        raise ValueError(f"Checksum set is not closed: {root}")
    actual = {
        path.relative_to(root).as_posix()
        for path in root.rglob("*")
        if path.is_file()
    }
    if actual != expected | {sums_name}:
        raise ValueError(
            f"File inventory is not closed: extra={actual - expected - {sums_name}}"
        )
    return manifest


def _labels(annotation: Mapping[str, Any]) -> tuple[Any, ...]:
    aspect_rows = annotation.get("aspects")
    if not isinstance(aspect_rows, list):
        raise ValueError("Malformed annotation aspects")
    by_aspect = {row.get("aspect"): row.get("label") for row in aspect_rows}
    if set(by_aspect) != set(ASPECTS):
        raise ValueError("Annotation aspect set mismatch")
    return tuple(by_aspect[aspect] for aspect in ASPECTS)


def validate(
    *,
    package: Path,
    source_release: Path,
    source_dir: Path,
    reference_package: Path,
) -> dict[str, Any]:
    try:
        import openpyxl
    except ModuleNotFoundError as exc:
        raise RuntimeError("openpyxl is required for XLSX validation") from exc

    package = package.resolve()
    source_release = source_release.resolve()
    source_dir = source_dir.resolve()
    reference_package = reference_package.resolve()
    standard = validate_release(
        package=package,
        release=package / "final",
        expected_records=9772,
    )
    source_manifest = _validate_closure(source_release)
    projection_root = package / "legacy_projection"
    projection_manifest = _validate_closure(projection_root)

    clean_rows = _read_jsonl(source_release / "clean_core.jsonl")
    source_ledger = _read_jsonl(
        source_release / "source_row_decision_ledger.jsonl"
    )
    blind_rows = _read_jsonl(package / "input" / "blind_reviews.jsonl")
    if (
        len(clean_rows) != 9772
        or len(blind_rows) != 9772
        or len(source_ledger) != 10105
    ):
        raise ValueError("Legacy source/prepared row counts do not close")
    if len({row["sample_id"] for row in clean_rows}) != 9772:
        raise ValueError("Clean-core sample IDs are not unique")
    if len({row["review_text_sha256"] for row in blind_rows}) != 9772:
        raise ValueError("Blind target text hashes are not unique")
    forbidden = set(ASPECTS)
    if any(forbidden.intersection(row) for row in clean_rows + blind_rows):
        raise ValueError("Historical aspect columns leaked into source/targets")
    scrub = _read_json(source_release / "legacy_label_scrub_audit.json")
    if (
        scrub.get("status") != "PASS"
        or scrub.get("historical_label_values_in_clean_core") != 0
        or scrub.get("historical_label_values_in_llm_target_input") != 0
        or scrub.get("source_workbooks_modified") is not False
    ):
        raise ValueError("Legacy label scrub audit failed")

    inventory = _read_json(
        source_release / "source_workbook_inventory.json"
    )
    for item in inventory["workbooks"]:
        source_path = REPOSITORY_ROOT / item["path"]
        if _sha256_file(source_path) != item["sha256"]:
            raise ValueError(f"Historical workbook changed: {source_path}")

    new_manifest = _read_json(package / "final" / "manifest.json")
    reference_manifest = _read_json(
        reference_package / "final" / "manifest.json"
    )
    new_primary_config = new_manifest["primary_run"]["execution_config"]
    new_diagnostic_config = new_manifest["diagnostic_run"]["execution_config"]
    reference_primary_config = reference_manifest["primary_run"][
        "execution_config"
    ]
    if new_primary_config != reference_primary_config:
        raise ValueError("Primary execution config differs from prior tranche")
    if new_diagnostic_config != reference_primary_config:
        raise ValueError("Diagnostic execution config differs from prior tranche")

    pseudo_rows = _read_jsonl(
        package / "final" / "ai_pseudo_labels.jsonl"
    )
    pseudo_by_sample = {row["sample_id"]: row for row in pseudo_rows}
    expanded_rows = _read_jsonl(
        projection_root / "legacy_old_relabel_10105.jsonl"
    )
    if len(pseudo_by_sample) != 9772 or len(expanded_rows) != 10105:
        raise ValueError("Canonical/expanded publication counts do not close")
    for row in expanded_rows:
        if row.get("legacy_label_values_reused") is not False:
            raise ValueError("Expanded row claims historical-label reuse")
        pseudo = pseudo_by_sample.get(row["canonical_sample_id"])
        if pseudo is None:
            raise ValueError("Expanded row lacks canonical pseudo-label")
        if (
            row["annotation_id"] != pseudo["annotation_id"]
            or _labels(row["annotation"]) != _labels(pseudo["annotation"])
        ):
            raise ValueError("Expanded annotation does not match canonical row")

    expanded_by_file: dict[str, list[dict[str, Any]]] = {}
    for row in expanded_rows:
        expanded_by_file.setdefault(row["source_file"], []).append(row)
    for rows in expanded_by_file.values():
        rows.sort(key=lambda item: item["source_row_number"])
    xlsx_rows = 0
    for xlsx_path in sorted(
        (projection_root / "xlsx_10col_new_labels").glob("*.xlsx")
    ):
        source_name = xlsx_path.name.replace(
            "_relabel_v2.xlsx",
            "_labeled.xlsx",
        )
        workbook = openpyxl.load_workbook(
            xlsx_path,
            read_only=True,
            data_only=True,
        )
        worksheet = workbook.active
        headers = tuple(
            cell.value
            for cell in next(worksheet.iter_rows(min_row=1, max_row=1))
        )
        if headers != EXPECTED_HEADERS:
            workbook.close()
            raise ValueError(f"Projected XLSX header mismatch: {xlsx_path}")
        expected_file_rows = expanded_by_file.get(source_name, [])
        projected_values = list(
            worksheet.iter_rows(min_row=2, values_only=True)
        )
        if len(projected_values) != len(expected_file_rows):
            workbook.close()
            raise ValueError(
                f"Projected XLSX row count mismatch: {xlsx_path}"
            )
        for values, expanded in zip(
            projected_values,
            expected_file_rows,
            strict=True,
        ):
            if expanded["source_sheet"] != worksheet.title:
                workbook.close()
                raise ValueError(
                    f"Projected XLSX sheet mismatch: {expanded['source_position']}"
                )
            if tuple(values[1:]) != _labels(expanded["annotation"]):
                workbook.close()
                raise ValueError(
                    f"Projected XLSX labels mismatch: "
                    f"{expanded['source_position']}"
                )
            xlsx_rows += 1
        workbook.close()
    if xlsx_rows != 10105:
        raise ValueError(f"Projected XLSX count mismatch: {xlsx_rows}")

    return {
        "status": "VALID",
        "source_rows": 10105,
        "canonical_unique_records": 9772,
        "duplicate_alias_rows": 333,
        "expanded_projection_rows": len(expanded_rows),
        "projected_xlsx_rows": xlsx_rows,
        "historical_label_values_in_llm_targets": 0,
        "historical_source_workbooks_modified": False,
        "primary_config_exact_match_prior_tranche": True,
        "diagnostic_config_exact_match_prior_tranche": True,
        "standard_release_validation": standard,
        "source_release_manifest_sha256": _sha256_file(
            source_release / "manifest.json"
        ),
        "canonical_final_manifest_sha256": _sha256_file(
            package / "final" / "manifest.json"
        ),
        "projection_manifest_sha256": _sha256_file(
            projection_root / "manifest.json"
        ),
        "cross_corpus_reserved_text_hash_overlap": (
            _read_json(package / "prepare_manifest.json")["selection"][
                "cross_corpus_reserved_text_hash_overlap"
            ]
        ),
        "interpretation": (
            "Schema/provenance-valid AI pseudo-labels pending human "
            "verification; this is not a human-gold accuracy claim."
        ),
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser()
    parser.add_argument("--package", type=Path, default=DEFAULT_PACKAGE)
    parser.add_argument(
        "--source-release",
        type=Path,
        default=DEFAULT_SOURCE_RELEASE,
    )
    parser.add_argument("--source-dir", type=Path, default=DEFAULT_SOURCE_DIR)
    parser.add_argument(
        "--reference-package",
        type=Path,
        default=DEFAULT_REFERENCE_PACKAGE,
    )
    parser.add_argument("--report", type=Path, default=DEFAULT_REPORT)
    return parser


def main() -> None:
    args = build_parser().parse_args()
    result = validate(
        package=args.package,
        source_release=args.source_release,
        source_dir=args.source_dir,
        reference_package=args.reference_package,
    )
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(
        json.dumps(result, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
        newline="\n",
    )
    print(json.dumps(result, ensure_ascii=False, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
