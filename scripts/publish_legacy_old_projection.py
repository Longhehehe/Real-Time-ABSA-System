"""Publish the new pseudo-labels back onto all 10,105 legacy source rows.

The standard ABSA release remains canonical at one row per normalized-unique
review.  This companion projection restores the original ten-workbook row
layout, maps every duplicate alias to its canonical new annotation, and writes
new XLSX files.  It never edits the historical workbooks in place.
"""

from __future__ import annotations

import argparse
from collections import Counter
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import shutil
import tempfile
from typing import Any


REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
ASPECTS = (
    "Chất lượng sản phẩm",
    "Hiệu năng & Trải nghiệm",
    "Đúng mô tả",
    "Giá cả & Khuyến mãi",
    "Vận chuyển",
    "Đóng gói",
    "Dịch vụ & Thái độ Shop",
    "Bảo hành & Đổi trả",
    "Tính xác thực",
)
EXPECTED_HEADERS = ("reviewContent", *ASPECTS)
DEFAULT_PACKAGE = Path(
    "data/annotations/absa_legacy_old_relabel_9772_v1_20260728"
)
DEFAULT_SOURCE_RELEASE = Path(
    "data/releases/legacy_old_reviews_label_blind_v1_20260728"
)
DEFAULT_SOURCE_DIR = Path("legacy/data/old_dataset")
DEFAULT_OUTPUT = DEFAULT_PACKAGE / "legacy_projection"


def _canonical_json(value: Any) -> str:
    return json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    )


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _read_json(path: Path) -> dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError(f"Expected JSON object: {path}")
    return value


def _read_jsonl(path: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    with path.open("r", encoding="utf-8") as handle:
        for line_number, line in enumerate(handle, 1):
            if not line.strip():
                continue
            value = json.loads(line)
            if not isinstance(value, dict):
                raise ValueError(f"Expected object at {path}:{line_number}")
            rows.append(value)
    return rows


def _write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
        newline="\n",
    )


def _write_jsonl(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="\n") as handle:
        for row in rows:
            handle.write(_canonical_json(row) + "\n")


def _artifact(
    path: Path,
    root: Path,
    *,
    records: int | None = None,
) -> dict[str, Any]:
    item: dict[str, Any] = {
        "path": path.relative_to(root).as_posix(),
        "bytes": path.stat().st_size,
        "sha256": _sha256_file(path),
    }
    if records is not None:
        item["records"] = records
    return item


def _natural_part(path: Path) -> tuple[int, str]:
    match = re.search(r"part(\d+)", path.stem, flags=re.IGNORECASE)
    return (int(match.group(1)) if match else 10**9, path.name.casefold())


def _labels(annotation: dict[str, Any]) -> dict[str, Any]:
    rows = annotation.get("aspects")
    if not isinstance(rows, list):
        raise ValueError("Annotation aspects are malformed")
    by_aspect = {row.get("aspect"): row.get("label") for row in rows}
    if set(by_aspect) != set(ASPECTS):
        raise ValueError("Annotation aspect set is not exact")
    return by_aspect


def publish(
    *,
    package: Path,
    source_release: Path,
    source_dir: Path,
    output: Path,
) -> dict[str, Any]:
    try:
        import openpyxl
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "openpyxl is required for the legacy XLSX projection"
        ) from exc

    package = package.resolve()
    source_release = source_release.resolve()
    source_dir = source_dir.resolve()
    output = output.resolve()
    if output.exists():
        raise FileExistsError(f"Projection already exists: {output}")
    final_root = package / "final"
    final_manifest = _read_json(final_root / "manifest.json")
    if (
        final_manifest.get("status")
        != "AI_PSEUDO_LABEL_PENDING_HUMAN_VERIFICATION"
    ):
        raise ValueError("Canonical pseudo-label release is not publishable")

    pseudo_rows = _read_jsonl(final_root / "ai_pseudo_labels.jsonl")
    pseudo_by_sample = {row["sample_id"]: row for row in pseudo_rows}
    if len(pseudo_by_sample) != 9772 or len(pseudo_rows) != 9772:
        raise ValueError("Expected exactly 9,772 unique pseudo-labels")
    source_ledger = _read_jsonl(
        source_release / "source_row_decision_ledger.jsonl"
    )
    if len(source_ledger) != 10105:
        raise ValueError("Expected exactly 10,105 source-row decisions")
    if not {
        row["canonical_sample_id"] for row in source_ledger
    }.issubset(pseudo_by_sample):
        raise ValueError("Source aliases do not close over pseudo-label release")

    ledger_by_position = {
        row["source_position"]: row for row in source_ledger
    }
    if len(ledger_by_position) != len(source_ledger):
        raise ValueError("Duplicate source positions in decision ledger")
    source_inventory = _read_json(
        source_release / "source_workbook_inventory.json"
    )
    expected_sha = {
        Path(row["path"]).name: row["sha256"]
        for row in source_inventory["workbooks"]
    }

    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = Path(
        tempfile.mkdtemp(prefix=f".{output.name}.", dir=output.parent)
    )
    try:
        expanded_rows: list[dict[str, Any]] = []
        xlsx_artifacts: list[dict[str, Any]] = []
        status_counts: Counter[str] = Counter()
        decision_counts: Counter[str] = Counter()
        source_files = sorted(source_dir.glob("*.xlsx"), key=_natural_part)
        for source_path in source_files:
            if _sha256_file(source_path) != expected_sha.get(source_path.name):
                raise ValueError(
                    f"Historical workbook changed: {source_path.name}"
                )
            source_book = openpyxl.load_workbook(
                source_path,
                read_only=True,
                data_only=True,
            )
            source_sheet = source_book.active
            headers = tuple(
                cell.value
                for cell in next(
                    source_sheet.iter_rows(min_row=1, max_row=1)
                )
            )
            if headers != EXPECTED_HEADERS:
                source_book.close()
                raise ValueError(f"Header mismatch: {source_path.name}")

            projected_book = openpyxl.Workbook()
            projected_sheet = projected_book.active
            projected_sheet.title = source_sheet.title
            projected_sheet.append(EXPECTED_HEADERS)
            projected_records = 0
            for row_number, values in enumerate(
                source_sheet.iter_rows(min_row=2, values_only=True),
                2,
            ):
                raw_review = values[0]
                if raw_review is None or not str(raw_review).strip():
                    continue
                position = (
                    f"{source_path.name}:{source_sheet.title}:{row_number}"
                )
                source_row = ledger_by_position.get(position)
                if source_row is None:
                    raise ValueError(f"Missing source decision: {position}")
                pseudo = pseudo_by_sample[source_row["canonical_sample_id"]]
                aspect_labels = _labels(pseudo["annotation"])
                projected_sheet.append(
                    [str(raw_review)]
                    + [aspect_labels[aspect] for aspect in ASPECTS]
                )
                status = pseudo["annotation"]["annotation_status"]
                status_counts[status] += 1
                decision_counts[source_row["decision"]] += 1
                expanded_rows.append(
                    {
                        "schema_version": (
                            "legacy-old-expanded-pseudo-label/1.0.0"
                        ),
                        "source_file": source_path.name,
                        "source_sheet": source_sheet.title,
                        "source_row_number": row_number,
                        "source_position": position,
                        "source_decision": source_row["decision"],
                        "canonical_sample_id": pseudo["sample_id"],
                        "annotation_id": pseudo["annotation_id"],
                        "reviewContent": str(raw_review),
                        "canonical_reviewContent": pseudo["reviewContent"],
                        "annotation": pseudo["annotation"],
                        "artifact_status": pseudo["artifact_status"],
                        "human_verification_status": pseudo[
                            "human_verification_status"
                        ],
                        "legacy_label_values_reused": False,
                        "labels_origin": (
                            "NEW_GPT_5_6_TERRA_PSEUDO_LABEL_CANONICAL_MAPPING"
                        ),
                    }
                )
                projected_records += 1
            source_book.close()
            target_path = (
                temporary
                / "xlsx_10col_new_labels"
                / source_path.name.replace(
                    "_labeled.xlsx",
                    "_relabel_v2.xlsx",
                )
            )
            target_path.parent.mkdir(parents=True, exist_ok=True)
            projected_book.save(target_path)
            projected_book.close()
            xlsx_artifacts.append(
                _artifact(
                    target_path,
                    temporary,
                    records=projected_records,
                )
            )

        if len(expanded_rows) != 10105:
            raise ValueError(
                f"Expanded projection is incomplete: {len(expanded_rows)}"
            )
        expanded_path = temporary / "legacy_old_relabel_10105.jsonl"
        _write_jsonl(expanded_path, expanded_rows)
        artifacts = [
            _artifact(expanded_path, temporary, records=len(expanded_rows)),
            *xlsx_artifacts,
        ]
        manifest = {
            "schema_version": "legacy-old-projection-manifest/1.0.0",
            "artifact_type": "LEGACY_OLD_NEW_PSEUDO_LABEL_PROJECTION",
            "status": "PUBLISHED_PENDING_HUMAN_VERIFICATION",
            "built_at": datetime.now(timezone.utc).isoformat(),
            "canonical_unique_records": len(pseudo_rows),
            "expanded_source_rows": len(expanded_rows),
            "duplicate_alias_rows": len(expanded_rows) - len(pseudo_rows),
            "source_workbooks_modified": False,
            "historical_label_values_reused": False,
            "xlsx_projection": {
                "columns": list(EXPECTED_HEADERS),
                "interpretation": (
                    "The nine aspect columns contain only new AI "
                    "pseudo-labels. JSONL remains canonical because XLSX "
                    "cannot preserve evidence spans and generation provenance."
                ),
            },
            "source_bindings": {
                "canonical_final_manifest_sha256": _sha256_file(
                    final_root / "manifest.json"
                ),
                "canonical_final_checksums_sha256": _sha256_file(
                    final_root / "SHA256SUMS.txt"
                ),
                "source_release_manifest_sha256": _sha256_file(
                    source_release / "manifest.json"
                ),
                "source_row_ledger_sha256": _sha256_file(
                    source_release / "source_row_decision_ledger.jsonl"
                ),
            },
            "annotation_status_counts_source_rows": dict(
                sorted(status_counts.items())
            ),
            "source_decision_counts": dict(sorted(decision_counts.items())),
            "limitations": [
                "All labels remain AI pseudo-labels pending human review.",
                "Duplicate aliases inherit the canonical review annotation.",
                "The 10-column XLSX projection omits evidence, uncertainty, "
                "status, and generation provenance; use JSONL for research.",
            ],
            "artifacts": sorted(artifacts, key=lambda item: item["path"]),
        }
        manifest_path = temporary / "manifest.json"
        _write_json(manifest_path, manifest)
        sums = [
            (item["sha256"], item["path"]) for item in artifacts
        ] + [(_sha256_file(manifest_path), "manifest.json")]
        (temporary / "SHA256SUMS.txt").write_text(
            "".join(
                f"{digest}  {relative}\n"
                for digest, relative in sorted(
                    sums,
                    key=lambda item: item[1],
                )
            ),
            encoding="utf-8",
            newline="\n",
        )
        temporary.replace(output)
    except Exception:
        shutil.rmtree(temporary, ignore_errors=True)
        raise
    return manifest


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser()
    parser.add_argument("--package", type=Path, default=DEFAULT_PACKAGE)
    parser.add_argument(
        "--source-release",
        type=Path,
        default=DEFAULT_SOURCE_RELEASE,
    )
    parser.add_argument("--source-dir", type=Path, default=DEFAULT_SOURCE_DIR)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser


def main() -> None:
    args = build_parser().parse_args()
    result = publish(
        package=args.package,
        source_release=args.source_release,
        source_dir=args.source_dir,
        output=args.output,
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
